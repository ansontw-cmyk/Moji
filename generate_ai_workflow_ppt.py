from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import copy

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# ── Color palette ──────────────────────────────────────────────
C_RED    = RGBColor(0xC0, 0x00, 0x00)   # FET red
C_DARK   = RGBColor(0x1A, 0x1A, 0x2E)  # near-black navy
C_NAVY   = RGBColor(0x16, 0x21, 0x3E)  # dark navy
C_BLUE   = RGBColor(0x0F, 0x3C, 0x78)  # mid blue
C_ACCENT = RGBColor(0x00, 0xB4, 0xD8)  # cyan accent
C_GOLD   = RGBColor(0xF5, 0xA6, 0x23)  # gold/amber
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_LGRAY  = RGBColor(0xF2, 0xF4, 0xF8)
C_MGRAY  = RGBColor(0x8A, 0x9B, 0xB0)
C_GREEN  = RGBColor(0x2D, 0xC6, 0x53)

BLANK = prs.slide_layouts[6]   # completely blank

# ── Helpers ────────────────────────────────────────────────────
def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=None):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.line.fill.background()
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        if line_w: shape.line.width = line_w
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, x, y, w, h, size=18, bold=False, color=C_WHITE,
             align=PP_ALIGN.LEFT, wrap=True, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.word_wrap = wrap
    tf = tb.text_frame
    tf.word_wrap = wrap
    p  = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return tb

def bg_dark(slide):
    add_rect(slide, 0, 0, 13.33, 7.5, fill=C_DARK)

def red_bar(slide, label="遠傳 FET", right_text="AI 工作流"):
    add_rect(slide, 0, 6.9, 13.33, 0.6, fill=C_RED)
    add_text(slide, label, 0.2, 6.92, 3, 0.45, size=11, bold=True, color=C_WHITE)
    add_text(slide, right_text, 10, 6.92, 3, 0.45, size=10, color=C_WHITE, align=PP_ALIGN.RIGHT)

def slide_num(slide, n):
    add_text(slide, str(n), 12.9, 6.92, 0.4, 0.4, size=11, color=C_WHITE, align=PP_ALIGN.RIGHT)

def accent_line(slide, x, y, w, color=C_RED):
    add_rect(slide, x, y, w, 0.04, fill=color)

def tag(slide, text, x, y, bg=C_RED, fg=C_WHITE, size=11):
    add_rect(slide, x, y, len(text)*0.13+0.3, 0.32, fill=bg)
    add_text(slide, text, x+0.1, y+0.02, len(text)*0.13+0.2, 0.3, size=size, bold=True, color=fg)


# ══════════════════════════════════════════════════════════════════
# Slide 1 — Title / 30分鐘震撼彈
# ══════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
bg_dark(s1)
# Left half navy panel
add_rect(s1, 0, 0, 6.5, 6.9, fill=C_NAVY)
# Red accent stripe
add_rect(s1, 0, 0, 0.18, 6.9, fill=C_RED)

# Part label
tag(s1, "PART 2", 0.5, 0.35, bg=C_RED)
add_text(s1, "AI 工作流", 0.5, 0.85, 5.8, 1.1, size=46, bold=True, color=C_WHITE)
accent_line(s1, 0.5, 2.0, 4.5, color=C_RED)

add_text(s1, "30 分鐘", 0.5, 2.15, 5.5, 1.0, size=52, bold=True, color=C_GOLD)
add_text(s1, "從產品構想到可執行程式碼", 0.5, 3.25, 5.5, 0.7, size=22, bold=False, color=C_ACCENT)

add_text(s1, "這不是 ChatGPT 的進階用法\n這是一支 AI 工程團隊同步運作",
         0.5, 4.1, 5.5, 1.2, size=15, color=RGBColor(0xB0,0xC4,0xDE))

# Right half — stat boxes
boxes = [
    ("7", "個 Agent 角色", C_RED),
    ("30", "分鐘完成", C_ACCENT),
    ("0", "次人工交接", C_GREEN),
]
for i,(num,label,col) in enumerate(boxes):
    bx = 7.1; by = 0.8 + i*1.9
    add_rect(s1, bx, by, 5.8, 1.6, fill=C_NAVY)
    add_rect(s1, bx, by, 0.12, 1.6, fill=col)
    add_text(s1, num,   bx+0.35, by+0.05, 2.5, 1.1, size=62, bold=True, color=col)
    add_text(s1, label, bx+0.35, by+1.1,  5.0, 0.45, size=16, color=C_WHITE)

add_text(s1, "vs. 傳統流程需 2–4 週", 7.1, 6.45, 5.8, 0.4,
         size=13, italic=True, color=RGBColor(0x88,0xAA,0xCC))

red_bar(s1); slide_num(s1, 1)


# ══════════════════════════════════════════════════════════════════
# Slide 2 — 單一 Agent 天花板
# ══════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
bg_dark(s2)
add_rect(s2, 0, 0, 13.33, 1.15, fill=C_NAVY)
add_rect(s2, 0, 0, 0.18, 7.5, fill=C_RED)

tag(s2, "問題診斷", 0.4, 0.15, bg=C_RED)
add_text(s2, "單打獨鬥的 AI，天花板在這裡", 0.4, 0.55, 12, 0.65, size=28, bold=True)

# Three problem cards
problems = [
    ("01", "Context Window\n有限", "複雜任務塞不進\n單一對話視窗", C_RED),
    ("02", "角色衝突", "既當 PM 又當工程師\n品質被稀釋", C_GOLD),
    ("03", "無法並行", "所有思考序列化\n效率卡在線性速度", C_ACCENT),
]
for i,(num,title,desc,col) in enumerate(problems):
    cx = 0.6 + i*4.15; cy = 1.4
    add_rect(s2, cx, cy, 3.85, 4.2, fill=C_NAVY)
    add_rect(s2, cx, cy, 3.85, 0.08, fill=col)
    add_text(s2, num,   cx+0.2, cy+0.18, 1.5, 0.8, size=42, bold=True, color=col)
    add_text(s2, title, cx+0.2, cy+1.05, 3.4, 0.9, size=20, bold=True, color=C_WHITE)
    accent_line(s2, cx+0.2, cy+2.0, 3.2, color=col)
    add_text(s2, desc,  cx+0.2, cy+2.15, 3.4, 1.7, size=15, color=RGBColor(0xB0,0xC4,0xDE))

# Solution call-out
add_rect(s2, 0.4, 5.8, 12.4, 0.75, fill=RGBColor(0x0F,0x3C,0x78))
add_text(s2, "解法：讓每個 Agent 專注單一角色，透過結構化 Output 串聯 → Multi-Agent Workflow",
         0.6, 5.88, 12.0, 0.55, size=15, bold=True, color=C_ACCENT)

red_bar(s2, right_text="AI 工作流"); slide_num(s2, 2)


# ══════════════════════════════════════════════════════════════════
# Slide 3 — 三大設計原則
# ══════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
bg_dark(s3)
add_rect(s3, 0, 0, 13.33, 1.15, fill=C_NAVY)
add_rect(s3, 0, 0, 0.18, 7.5, fill=C_RED)

tag(s3, "設計原則", 0.4, 0.15, bg=C_RED)
add_text(s3, "Multi-Agent 三大核心設計", 0.4, 0.55, 12, 0.65, size=28, bold=True)

principles = [
    ("角色專一化\nRole Specialization",
     "每個 Agent 只做一件事\nPrompt 設計 = JD 職位說明書\n\n✦  PM Agent 只輸出 PRD\n✦  UX Agent 只輸出線框圖規格\n✦  QA Agent 只做驗證",
     C_RED),
    ("結構化交接\nStructured Handoff",
     "Output 必須是下一個 Agent\n的有效 Input\n\n✦  格式：Markdown / JSON Schema\n✦  工程師設計「格式」\n   不是「對話內容」\n✦  錯誤在格式邊界被攔截",
     C_ACCENT),
    ("編排策略\nOrchestration Pattern",
     "Sequential  A→B→C\n（有依賴關係）\n\nParallel   A╮\n           B╯→C\n（獨立子任務同時跑）\n\nConditional  依結果走不同路",
     C_GOLD),
]
for i,(title,body,col) in enumerate(principles):
    cx = 0.5 + i*4.25; cy = 1.35
    add_rect(s3, cx, cy, 3.95, 5.1, fill=C_NAVY)
    add_rect(s3, cx, cy, 0.1, 5.1, fill=col)
    add_text(s3, f"0{i+1}", cx+0.3, cy+0.12, 1.5, 0.65, size=32, bold=True, color=col)
    add_text(s3, title,    cx+0.3, cy+0.75, 3.5, 0.9, size=15, bold=True, color=C_WHITE)
    accent_line(s3, cx+0.3, cy+1.65, 3.3, color=col)
    add_text(s3, body,     cx+0.3, cy+1.8,  3.5, 3.1, size=12.5, color=RGBColor(0xB8,0xCC,0xE0))

red_bar(s3, right_text="AI 工作流"); slide_num(s3, 3)


# ══════════════════════════════════════════════════════════════════
# Slide 4 — 案例深潛（Pipeline 圖）
# ══════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(BLANK)
bg_dark(s4)
add_rect(s4, 0, 0, 13.33, 1.1, fill=C_NAVY)
add_rect(s4, 0, 0, 0.18, 7.5, fill=C_RED)

tag(s4, "真實案例", 0.4, 0.14, bg=C_RED)
add_text(s4, "遠傳心生活 Super App — 7 Agents，產品從零到一", 0.4, 0.52, 12.5, 0.62, size=24, bold=True)

# Stat strip
for i,(val,lab) in enumerate([("7","Agents"),("30 min","完成"),("0","人工交接")]):
    sx = 1.5 + i*3.8
    add_rect(s4, sx, 1.18, 3.3, 0.52, fill=C_BLUE)
    add_text(s4, val, sx+0.15, 1.2, 1.6, 0.46, size=20, bold=True, color=C_GOLD)
    add_text(s4, lab, sx+1.4,  1.26, 1.7, 0.38, size=13, color=C_WHITE)

# Pipeline phases
# Phase labels
add_rect(s4, 0.3, 1.85, 1.55, 1.55, fill=RGBColor(0x8B,0x00,0x00))
add_text(s4, "PHASE\n1\n並行", 0.35, 1.9, 1.4, 1.4, size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

add_rect(s4, 0.3, 3.55, 1.55, 1.0, fill=C_BLUE)
add_text(s4, "PHASE\n2", 0.35, 3.6, 1.4, 0.9, size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

add_rect(s4, 0.3, 4.68, 1.55, 0.9, fill=C_BLUE)
add_text(s4, "PHASE\n3", 0.35, 4.73, 1.4, 0.8, size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

add_rect(s4, 0.3, 5.7, 1.55, 1.55, fill=RGBColor(0x8B,0x00,0x00))
add_text(s4, "PHASE\n4\n並行", 0.35, 5.75, 1.4, 1.4, size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

# Agent boxes
agents = [
    (2.1, 1.88, 4.7, 0.68, "Trend Researcher", "市場趨勢 + 競品矩陣", C_GOLD),
    (2.1, 2.68, 4.7, 0.68, "Product Manager",  "PRD + User Stories + RICE", C_GOLD),
    (2.1, 3.58, 4.7, 0.82, "UX Architect",      "用戶旅程 + 線框圖規格", C_ACCENT),
    (2.1, 4.55, 4.7, 0.88, "Software Architect","系統架構 + API Contract + DB Schema", C_RED),
    (2.1, 5.55, 4.7, 0.68, "Frontend Developer","HTML/CSS/JS 原型", C_GREEN),
    (2.1, 6.32, 4.7, 0.68, "Backend Architect", "Node.js API + DB + 認證", C_GREEN),
]
for (bx,by,bw,bh,name,desc,col) in agents:
    add_rect(s4, bx, by, bw, bh, fill=C_NAVY)
    add_rect(s4, bx, by, 0.09, bh, fill=col)
    add_text(s4, name, bx+0.2, by+0.04, bw-0.3, 0.35, size=13.5, bold=True, color=col)
    add_text(s4, desc, bx+0.2, by+0.36, bw-0.3, 0.38, size=11, color=RGBColor(0xB0,0xC4,0xDE))

# Arrows / connectors (simple text arrows)
arrow_y = [2.6, 3.45, 4.48, 5.52]
for ay in arrow_y:
    add_text(s4, "▼", 4.2, ay, 0.5, 0.28, size=13, color=C_RED, align=PP_ALIGN.CENTER)

# Right side: output boxes
outputs = [
    (7.1, 1.88, 5.8, 1.42, "Phase 1 輸出", "市場報告 + PRD + User Stories\nNorth Star Metric + RICE 優先級", C_GOLD),
    (7.1, 3.42, 5.8, 0.95, "Phase 2 輸出", "用戶旅程圖 + IA + 線框圖規格\n5 大互動設計原則", C_ACCENT),
    (7.1, 4.48, 5.8, 0.98, "Phase 3 輸出", "API Contract + DB Schema\n< 300ms SLA 架構 + 個資合規設計", C_RED),
    (7.1, 5.55, 5.8, 1.28, "Phase 4 輸出", "3 個可運作 HTML 頁面原型\n6 個 RESTful API endpoints 上線", C_GREEN),
]
for (ox,oy,ow,oh,title,body,col) in outputs:
    add_rect(s4, ox, oy, ow, oh, fill=C_NAVY)
    add_rect(s4, ox, oy, ow, 0.07, fill=col)
    add_text(s4, title, ox+0.15, oy+0.1,  ow-0.3, 0.35, size=12, bold=True, color=col)
    add_text(s4, body,  ox+0.15, oy+0.48, ow-0.3, oh-0.55, size=11.5, color=C_WHITE)

# Connector lines (dashes)
for oy in [2.22, 3.95, 4.98, 5.88]:
    add_text(s4, "─ ─ ─ ─ →", 6.85, oy, 0.6, 0.3, size=11, color=C_MGRAY)

# Reality Checker badge
add_rect(s4, 9.5, 1.0, 3.5, 0.5, fill=C_RED)
add_text(s4, "Phase 5 → Reality Checker QA 驗證", 9.55, 1.04, 3.4, 0.4,
         size=11, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

red_bar(s4, right_text="AI 工作流"); slide_num(s4, 4)


# ══════════════════════════════════════════════════════════════════
# Slide 5 — 技術底層 LLM + Skills + MCP
# ══════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(BLANK)
bg_dark(s5)
add_rect(s5, 0, 0, 13.33, 1.1, fill=C_NAVY)
add_rect(s5, 0, 0, 0.18, 7.5, fill=C_RED)

tag(s5, "技術底層", 0.4, 0.14, bg=C_RED)
add_text(s5, "成功關鍵不是模型有多強，而是 Orchestration 設計有多清晰", 0.4, 0.52, 12.5, 0.62, size=22, bold=True)

# Architecture stack (left)
layers = [
    ("Orchestrator 調度中心",  "決定誰先跑 / 誰等誰 / 錯誤如何處理",           C_RED,    0.5, 1.3,  6.0, 0.95),
    ("Agent 1…N  各司其職",    "每個 Agent = 專業角色 + 獨立 Context",         C_BLUE,   0.5, 2.4,  6.0, 0.95),
    ("Skills  專業知識",       "封裝可重用的能力模組（搜尋/分析/格式化輸出）",    C_ACCENT, 0.5, 3.5,  6.0, 0.95),
    ("MCP  手腳",              "連接外部系統（DB / API / 程式碼執行環境）",       C_GOLD,   0.5, 4.6,  6.0, 0.95),
    ("外部系統",               "現有子系統 API / 資料庫 / 第三方服務",           C_GREEN,  0.5, 5.7,  6.0, 0.8),
]
for (title,desc,col,lx,ly,lw,lh) in layers:
    add_rect(s5, lx, ly, lw, lh, fill=C_NAVY)
    add_rect(s5, lx, ly, 0.1, lh, fill=col)
    add_text(s5, title, lx+0.25, ly+0.08, lw-0.35, 0.42, size=16, bold=True, color=col)
    add_text(s5, desc,  lx+0.25, ly+0.48, lw-0.35, 0.42, size=12, color=RGBColor(0xB0,0xC4,0xDE))

# Arrows between layers
for ay in [2.28, 3.38, 4.48, 5.58]:
    add_text(s5, "↓", 3.4, ay, 0.4, 0.2, size=14, color=C_RED, align=PP_ALIGN.CENTER)

# Right panel — connection to Part 1
add_rect(s5, 7.0, 1.3, 5.9, 5.2, fill=C_NAVY)
add_rect(s5, 7.0, 1.3, 5.9, 0.07, fill=C_RED)
add_text(s5, "與第一部分的連結", 7.15, 1.38, 5.6, 0.48, size=16, bold=True, color=C_RED)
add_text(s5,
    "第一部分建立的 Skills + MCP\n就是讓這些 Agent 真正能動的基礎建設\n\n"
    "Skills（專業知識）\n"
    "  ✦  Requirement AI Auditor\n"
    "  ✦  Security & Logic Auditor\n"
    "  ✦  Auto Testing Agent\n"
    "  ✦  Diagnostic & Analysis Agent\n\n"
    "MCP（手腳）\n"
    "  ✦  10+ MCP Services 已上線\n"
    "  ✦  連接程式碼執行 / DB / 外部 API\n\n"
    "→ 社群 20+ Skills 可直接被 Workflow 呼叫",
    7.15, 1.95, 5.6, 4.4, size=13, color=C_WHITE)

# Highlight callout
add_rect(s5, 7.0, 6.0, 5.9, 0.52, fill=RGBColor(0x0F,0x3C,0x78))
add_text(s5, "AI 工作流 = 第一部分積木的組裝方式",
         7.15, 6.05, 5.7, 0.42, size=14, bold=True, color=C_ACCENT)

red_bar(s5, right_text="AI 工作流"); slide_num(s5, 5)


# ══════════════════════════════════════════════════════════════════
# Slide 6 — 你的場景怎麼用
# ══════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(BLANK)
bg_dark(s6)
add_rect(s6, 0, 0, 13.33, 1.1, fill=C_NAVY)
add_rect(s6, 0, 0, 0.18, 7.5, fill=C_RED)

tag(s6, "應用場景", 0.4, 0.14, bg=C_RED)
add_text(s6, "任何「多角色、多步驟、有依賴關係」的工作都可以 Workflow 化", 0.4, 0.52, 12.5, 0.62, size=20, bold=True)

scenarios = [
    ("客服升級",
     "意圖分析 → 知識檢索\n→ 回應生成 → 品質審核",
     "處理時間 -60%", C_ACCENT),
    ("網路故障診斷",
     "告警收集 → 根因分析\n→ 修復建議 → 工單生成",
     "MTTR 大幅縮短", C_RED),
    ("行銷內容工廠",
     "市場研究 → 文案生成\n→ SEO 審核 → 多平台輸出",
     "產出量 10×", C_GOLD),
    ("法規合規審查",
     "文件解析 → 條文比對\n→ 風險標記 → 報告生成",
     "人工審查 -80%", C_GREEN),
]
for i,(title,flow,result,col) in enumerate(scenarios):
    col_idx = i % 2; row_idx = i // 2
    cx = 0.4 + col_idx*6.5; cy = 1.45 + row_idx*2.5
    add_rect(s6, cx, cy, 6.1, 2.2, fill=C_NAVY)
    add_rect(s6, cx, cy, 6.1, 0.07, fill=col)
    add_rect(s6, cx, cy, 0.1, 2.2, fill=col)
    add_text(s6, title,  cx+0.25, cy+0.1,  5.7, 0.48, size=18, bold=True, color=col)
    add_text(s6, flow,   cx+0.25, cy+0.65, 5.7, 0.95, size=13, color=C_WHITE)
    # Result badge
    add_rect(s6, cx+3.5, cy+1.72, 2.4, 0.38, fill=col)
    add_text(s6, result, cx+3.55, cy+1.75, 2.3, 0.32, size=13, bold=True, color=C_WHITE)

# Interactive callout
add_rect(s6, 0.4, 6.45, 12.5, 0.45, fill=RGBColor(0x0F,0x3C,0x78))
add_text(s6, "💬  你現在浮現哪個場景？邀請 1-2 位說出來，現場快速拆解 Agent 架構",
         0.6, 6.48, 12.1, 0.38, size=13.5, bold=True, color=C_ACCENT)

red_bar(s6, right_text="AI 工作流"); slide_num(s6, 6)


# ══════════════════════════════════════════════════════════════════
# Slide 7 — 結語 + 行動呼籲
# ══════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(BLANK)
bg_dark(s7)
add_rect(s7, 0, 0, 13.33, 7.5, fill=C_DARK)
add_rect(s7, 0, 0, 0.18, 7.5, fill=C_RED)

add_text(s7, "AI 工作流不是未來", 0.45, 0.4, 12.5, 0.85, size=38, bold=True, color=C_WHITE)
add_text(s7, "是現在可以開始的事", 0.45, 1.2, 12.5, 0.85, size=38, bold=True, color=C_GOLD)
accent_line(s7, 0.45, 2.1, 6.0, color=C_RED)

# 3 Steps
steps = [
    ("01", "識別場景",
     "找出工作中一個需要\n3 個不同專業視角的任務\n這就是你的第一個 Multi-Agent 場景",
     C_RED),
    ("02", "套用框架",
     "角色專一 + 結構化交接\n選好 Sequential / Parallel 編排\nOuput 格式設計好，其餘交給 AI",
     C_ACCENT),
    ("03", "加入社群",
     "遠傳 AI 工作室\n20+ Skills、10+ MCP 可直接取用\n不用從零開始",
     C_GREEN),
]
for i,(num,title,body,col) in enumerate(steps):
    sx = 0.5 + i*4.2; sy = 2.3
    add_rect(s7, sx, sy, 3.9, 4.0, fill=C_NAVY)
    add_rect(s7, sx, sy, 3.9, 0.07, fill=col)
    add_text(s7, num,   sx+0.2, sy+0.15, 1.2, 0.75, size=36, bold=True, color=col)
    add_text(s7, title, sx+0.2, sy+0.95, 3.5, 0.55, size=18, bold=True, color=C_WHITE)
    accent_line(s7, sx+0.2, sy+1.55, 3.3, color=col)
    add_text(s7, body,  sx+0.2, sy+1.72, 3.5, 2.1, size=13, color=RGBColor(0xB0,0xC4,0xDE))

# Final CTA
add_rect(s7, 0.45, 6.45, 12.5, 0.45, fill=C_RED)
add_text(s7, "下週就可以有產出 — 從 3 個 Agent 的小流程開始",
         0.65, 6.48, 12.0, 0.38, size=15, bold=True, color=C_WHITE)

red_bar(s7, right_text="AI 工作流"); slide_num(s7, 7)


# ── Save ───────────────────────────────────────────────────────
out = "/home/user/Moji/AI_Workflow_Part2.pptx"
prs.save(out)
print(f"Saved: {out}")
