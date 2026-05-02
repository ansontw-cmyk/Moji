#!/usr/bin/env python3
"""Multi-Agent Databricks Intelligence Platform PPT Generator"""

import io
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import matplotlib.patheffects as pe

# ── Color Palette ──────────────────────────────────────────────────────────
NAVY    = RGBColor(0x0A, 0x16, 0x28)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BLUE    = RGBColor(0x00, 0xB4, 0xD8)
GREEN   = RGBColor(0x06, 0xD6, 0xA0)
ORANGE  = RGBColor(0xFF, 0x6B, 0x35)
PURPLE  = RGBColor(0x7B, 0x2F, 0xBE)
LGRAY   = RGBColor(0xF4, 0xF6, 0xF9)
DGRAY   = RGBColor(0x3D, 0x3D, 0x3D)
GOLD    = RGBColor(0xFF, 0xC3, 0x00)

W, H = Inches(13.33), Inches(7.5)   # 16:9

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]   # completely blank

# ── Helpers ────────────────────────────────────────────────────────────────

def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=Pt(0)):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.fill.background()
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = line_w
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, x, y, w, h, size=18, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(x, y, w, h)
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    return txb

def bg(slide, color=NAVY):
    add_rect(slide, 0, 0, W, H, fill=color)

def chart_img(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format='png', bbox_inches='tight', dpi=150,
                facecolor=fig.get_facecolor())
    buf.seek(0)
    plt.close(fig)
    return buf

def add_img(slide, buf, x, y, w, h):
    slide.shapes.add_picture(buf, x, y, w, h)

def accent_bar(slide, color=BLUE, y=Inches(1.15)):
    add_rect(slide, Inches(0.5), y, Inches(0.08), Inches(0.55), fill=color)

def slide_header(slide, title, subtitle='', title_color=NAVY,
                 accent_color=BLUE, light_bg=True):
    if light_bg:
        bg(slide, LGRAY)
        add_rect(slide, 0, 0, W, Inches(1.0), fill=WHITE)
    accent_bar(slide, accent_color, Inches(0.22))
    add_text(slide, title,    Inches(0.7), Inches(0.18), Inches(11), Inches(0.7),
             size=28, bold=True, color=title_color, align=PP_ALIGN.LEFT)
    if subtitle:
        add_text(slide, subtitle, Inches(0.7), Inches(0.85), Inches(11), Inches(0.35),
                 size=13, color=DGRAY, align=PP_ALIGN.LEFT)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
# gradient overlay strip
add_rect(s, 0, 0, Inches(5), H, fill=RGBColor(0x00, 0x1A, 0x3A))
# accent lines
add_rect(s, Inches(0.5), Inches(1.8), Inches(0.1), Inches(3.2), fill=BLUE)
add_rect(s, Inches(0.7), Inches(1.8), Inches(0.04), Inches(3.2), fill=GREEN)
# title
add_text(s, '多 Agent AI 智慧決策平台', Inches(0.9), Inches(1.9),
         Inches(11.5), Inches(1.2), size=44, bold=True, color=WHITE)
add_text(s, 'Multi-Agent AI Intelligence Decision Platform',
         Inches(0.9), Inches(3.1), Inches(11.5), Inches(0.6),
         size=20, color=BLUE)
add_text(s, '基於 Databricks 的跨域資料融合與商業洞察',
         Inches(0.9), Inches(3.75), Inches(11.5), Inches(0.6),
         size=18, color=RGBColor(0xCC, 0xDD, 0xFF))
# bottom bar
add_rect(s, 0, Inches(6.7), W, Inches(0.8), fill=RGBColor(0x00, 0x1A, 0x3A))
add_text(s, '2025  |  Powered by AI Multi-Agent Architecture  |  Databricks Lakehouse',
         Inches(0.5), Inches(6.75), Inches(12), Inches(0.4),
         size=11, color=RGBColor(0x88, 0xAA, 0xCC), align=PP_ALIGN.CENTER)
# data icons row
icons = ['🛒 購物', '📱 電信', '💳 DCB', '🚗 停車', '🛡️ 保險']
for i, ic in enumerate(icons):
    x = Inches(0.9 + i * 2.35)
    add_rect(s, x, Inches(5.5), Inches(2.1), Inches(0.9),
             fill=RGBColor(0x00, 0x2A, 0x4A), line=BLUE, line_w=Pt(1))
    add_text(s, ic, x + Inches(0.1), Inches(5.58), Inches(1.9), Inches(0.7),
             size=14, color=WHITE, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Executive Summary
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, '執行摘要', 'Executive Summary — Key Highlights', accent_color=BLUE)

cards = [
    (BLUE,   '🔗 5大異質資料源整合',
     '購物 / 電信 / DCB / 停車 / 保險\n統一接入 Databricks Lakehouse'),
    (GREEN,  '🤖 9個專業AI Agent協同',
     '各司其職的智慧代理人\n自動協調、並行分析、交叉驗證'),
    (ORANGE, '⚡ 即時洞察 + 預測分析',
     '毫秒級查詢 × 機器學習模型\n從描述→預測→規範性決策'),
    (PURPLE, '💰 可量化的商業價值',
     '行銷ROI +35% / 流失率 -28%\n詐欺損失 -42% / 決策速度 +60%'),
]
for i, (col, title, body) in enumerate(cards):
    x = Inches(0.4 + i * 3.1)
    add_rect(s, x, Inches(1.5), Inches(2.95), Inches(4.8), fill=WHITE,
             line=col, line_w=Pt(2))
    add_rect(s, x, Inches(1.5), Inches(2.95), Inches(0.55), fill=col)
    add_text(s, title, x + Inches(0.1), Inches(1.52),
             Inches(2.75), Inches(0.52), size=12, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, body,  x + Inches(0.15), Inches(2.15),
             Inches(2.65), Inches(3.9), size=12.5, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Challenge & Opportunity
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, '挑戰與機會', 'Challenge & Opportunity', accent_color=ORANGE)
# left panel
add_rect(s, Inches(0.4), Inches(1.4), Inches(5.9), Inches(5.7),
         fill=RGBColor(0xFF, 0xF0, 0xEC), line=ORANGE, line_w=Pt(1.5))
add_rect(s, Inches(0.4), Inches(1.4), Inches(5.9), Inches(0.5), fill=ORANGE)
add_text(s, '⚠️  現有痛點', Inches(0.5), Inches(1.42),
         Inches(5.7), Inches(0.45), size=14, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
pains = ['● 資料孤島：各系統資料無法互通',
         '● 人工分析：報表產出慢，時效性差',
         '● 決策滯後：洞察到行動之間落差大',
         '● 人才瓶頸：數據科學家資源有限',
         '● 機會流失：無法即時捕捉消費信號']
for i, p in enumerate(pains):
    add_text(s, p, Inches(0.6), Inches(2.05 + i * 0.82),
             Inches(5.5), Inches(0.7), size=13, color=DGRAY)

# right panel
add_rect(s, Inches(6.8), Inches(1.4), Inches(5.9), Inches(5.7),
         fill=RGBColor(0xEC, 0xFB, 0xF5), line=GREEN, line_w=Pt(1.5))
add_rect(s, Inches(6.8), Inches(1.4), Inches(5.9), Inches(0.5), fill=GREEN)
add_text(s, '✅  Multi-Agent 解方', Inches(6.9), Inches(1.42),
         Inches(5.7), Inches(0.45), size=14, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
opps = ['◆ 統一平台：Databricks 整合所有資料',
        '◆ 自動分析：Agent 7×24 持續運行',
        '◆ 即時決策：秒級洞察直達業務端',
        '◆ 規模智慧：AI 代替人力執行分析',
        '◆ 精準捕捉：360度客戶行為即時感知']
for i, o in enumerate(opps):
    add_text(s, o, Inches(7.0), Inches(2.05 + i * 0.82),
             Inches(5.5), Inches(0.7), size=13, color=DGRAY)

# center arrow
add_text(s, '→', Inches(6.2), Inches(3.8), Inches(0.6), Inches(0.8),
         size=40, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Data Asset Landscape
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
add_rect(s, 0, 0, W, Inches(1.1), fill=RGBColor(0x00, 0x1A, 0x3A))
add_rect(s, Inches(0.5), Inches(0.2), Inches(0.08), Inches(0.7), fill=BLUE)
add_text(s, 'Databricks 資料資產全景', Inches(0.7), Inches(0.18),
         Inches(11), Inches(0.75), size=28, bold=True, color=WHITE)

assets = [
    ('🛒', '購物資料', 'Shopping Data',
     '• 消費行為記錄\n• 品類偏好分析\n• 購買頻率/金額\n• 會員積點軌跡',
     BLUE),
    ('📱', '電信資料', 'Telecom Data',
     '• 通話/簡訊記錄\n• 上網行為數據\n• 基地台位置軌跡\n• 方案訂閱狀態',
     GREEN),
    ('💳', 'DCB交易', 'DCB Transactions',
     '• 電信代收帳單\n• 小額支付記錄\n• 訂閱服務明細\n• 跨平台交易流',
     ORANGE),
    ('🚗', '停車資料', 'Parking Data',
     '• 停車場進出記錄\n• 地理位置分佈\n• 停留時段分析\n• 鄰近商圈關聯',
     PURPLE),
    ('🛡️', '保險資料', 'Insurance Data',
     '• 保單持有資訊\n• 理賠申請記錄\n• 個人風險評分\n• 產品交叉持有',
     GOLD),
]
for i, (icon, tw, en, body, col) in enumerate(assets):
    x = Inches(0.3 + i * 2.55)
    add_rect(s, x, Inches(1.25), Inches(2.4), Inches(5.8),
             fill=RGBColor(0x00, 0x22, 0x40), line=col, line_w=Pt(1.5))
    add_rect(s, x, Inches(1.25), Inches(2.4), Inches(1.0), fill=col)
    add_text(s, icon, x, Inches(1.25), Inches(2.4), Inches(1.0),
             size=28, align=PP_ALIGN.CENTER, color=WHITE)
    add_text(s, tw, x + Inches(0.1), Inches(2.35),
             Inches(2.2), Inches(0.45), size=14, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, en, x + Inches(0.1), Inches(2.78),
             Inches(2.2), Inches(0.35), size=10,
             color=col, align=PP_ALIGN.CENTER)
    add_text(s, body, x + Inches(0.15), Inches(3.2),
             Inches(2.1), Inches(3.6), size=11.5,
             color=RGBColor(0xCC, 0xDD, 0xFF))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Multi-Agent Architecture (matplotlib diagram)
# ══════════════════════════════════════════════════════════════════════════
fig, ax = plt.subplots(figsize=(13, 6.5), facecolor='#0A1628')
ax.set_facecolor('#0A1628')
ax.set_xlim(0, 13); ax.set_ylim(0, 6.5)
ax.axis('off')

# Center hub
hub = plt.Circle((6.5, 3.2), 1.1, color='#003366', zorder=3)
ax.add_patch(hub)
hub2 = plt.Circle((6.5, 3.2), 1.1, fill=False,
                   edgecolor='#00B4D8', linewidth=2.5, zorder=4)
ax.add_patch(hub2)
ax.text(6.5, 3.35, 'Databricks\nLakehouse', ha='center', va='center',
        fontsize=11, fontweight='bold', color='white', zorder=5)

agents_info = [
    ('資料協調\nOrchestrator',  '#00B4D8', 6.5, 5.8),
    ('購物分析\nShopping',      '#06D6A0', 3.0, 5.2),
    ('電信分析\nTelecom',       '#06D6A0', 1.1, 3.2),
    ('DCB交易\nDCB Agent',      '#FF6B35', 3.0, 1.2),
    ('停車洞察\nParking',       '#FF6B35', 6.5, 0.5),
    ('保險風控\nInsurance',     '#7B2FBE', 10.0, 1.2),
    ('跨域融合\nFusion',        '#7B2FBE', 11.9, 3.2),
    ('商業推薦\nBiz Reco',      '#FFC300', 10.0, 5.2),
    ('報告生成\nReport Gen',    '#FFC300', 8.2, 5.8),
]
for label, col, ax_, ay_ in agents_info:
    c = plt.Circle((ax_, ay_), 0.72, color=col, alpha=0.25, zorder=2)
    ax.add_patch(c)
    c2 = plt.Circle((ax_, ay_), 0.72, fill=False,
                    edgecolor=col, linewidth=1.8, zorder=3)
    ax.add_patch(c2)
    ax.text(ax_, ay_, label, ha='center', va='center',
            fontsize=8.5, color='white', fontweight='bold', zorder=4)
    # draw line to center
    dx, dy = 6.5 - ax_, 3.2 - ay_
    dist = (dx**2 + dy**2)**0.5
    ux, uy = dx/dist, dy/dist
    ax.annotate('', xy=(6.5 - ux*1.12, 3.2 - uy*1.12),
                xytext=(ax_ + ux*0.74, ay_ + uy*0.74),
                arrowprops=dict(arrowstyle='->', color=col,
                                lw=1.4, alpha=0.7))

# Output bar
ax.add_patch(plt.Rectangle((1.5, -0.05), 10, 0.35,
             color='#00334D', zorder=2))
ax.text(6.5, 0.12, '▶  Business Output:  Dashboard  |  API  |  Alert  |  Report  |  Recommendation',
        ha='center', va='center', fontsize=9, color='#00B4D8')

buf5 = chart_img(fig)

s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
add_rect(s, 0, 0, W, Inches(0.9), fill=RGBColor(0x00, 0x1A, 0x3A))
add_text(s, '多Agent協作架構 | Multi-Agent Architecture',
         Inches(0.5), Inches(0.1), Inches(12), Inches(0.75),
         size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_img(s, buf5, Inches(0.15), Inches(0.9), Inches(13.0), Inches(6.4))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Agent Roles
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, 'Agent 角色分工', 'Agent Roles & Responsibilities', accent_color=PURPLE)

rows = [
    ('資料協調 Orchestrator', '全域資料', '排程、協調、任務分派', '執行計劃 + 狀態報告', BLUE),
    ('購物分析 Shopping',     '消費交易', 'RFM分群、籃子分析',   '客群標籤 + 推薦清單', GREEN),
    ('電信分析 Telecom',      '通聯資料', '流失預警、生活型態',  '預警名單 + 行為畫像', GREEN),
    ('DCB交易 Payment',       '支付記錄', '消費模式、能力評估',  '信用標籤 + 異常警示', ORANGE),
    ('停車洞察 Parking',      '位置資料', '熱力分析、商圈識別',  '選址報告 + 投放建議', ORANGE),
    ('保險風控 Insurance',    '保單理賠', '風險評分、交叉推薦',  '風險分層 + 產品清單', PURPLE),
    ('跨域融合 Fusion',       '多源整合', '360°畫像、信號融合',  '統一客戶畫像',        GOLD),
    ('商業推薦 Biz Reco',     '洞察結果', '決策生成、行動建議',  '商機清單 + 策略書',   GOLD),
    ('報告生成 Reporter',     '全部輸出', 'PPT/PDF/Dashboard',   '自動化報告交付',      BLUE),
]
# header row
cols_x = [Inches(0.35), Inches(2.4), Inches(4.7), Inches(7.2), Inches(10.1)]
col_w  = [Inches(2.0),  Inches(2.2), Inches(2.4),  Inches(2.8),  Inches(2.8)]
hdrs   = ['Agent', '輸入資料', '核心功能', '輸出成果', '顏色']
add_rect(s, Inches(0.35), Inches(1.35), Inches(12.6), Inches(0.42), fill=NAVY)
for j, h in enumerate(hdrs[:4]):
    add_text(s, h, cols_x[j] + Inches(0.05), Inches(1.38),
             col_w[j], Inches(0.35), size=11, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)

for i, (agent, inp, func, out, col) in enumerate(rows):
    y = Inches(1.77 + i * 0.55)
    fill = RGBColor(0xF4, 0xF6, 0xF9) if i % 2 == 0 else WHITE
    add_rect(s, Inches(0.35), y, Inches(12.6), Inches(0.52), fill=fill)
    add_rect(s, Inches(0.35), y, Inches(0.12), Inches(0.52), fill=col)
    for j, txt in enumerate([agent, inp, func, out]):
        add_text(s, txt, cols_x[j] + Inches(0.12), y + Inches(0.08),
                 col_w[j] - Inches(0.15), Inches(0.38),
                 size=10.5, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Shopping Intelligence (with bar chart)
# ══════════════════════════════════════════════════════════════════════════
fig, ax = plt.subplots(figsize=(5.5, 3.8), facecolor='white')
segments = ['高價值\n客群', '成長型\n客群', '潛力\n客群', '流失\n風險', '沉睡\n客群']
values   = [22, 31, 18, 14, 15]
colors_b = ['#06D6A0','#00B4D8','#7B2FBE','#FF6B35','#aaaaaa']
bars = ax.bar(segments, values, color=colors_b, edgecolor='white', linewidth=1.2)
for bar, v in zip(bars, values):
    ax.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 0.5,
            f'{v}%', ha='center', va='bottom', fontsize=10, fontweight='bold')
ax.set_ylabel('客戶佔比 %', fontsize=10)
ax.set_title('RFM 客戶分群分佈', fontsize=12, fontweight='bold', pad=8)
ax.set_ylim(0, 40); ax.grid(axis='y', alpha=0.3)
ax.spines['top'].set_visible(False); ax.spines['right'].set_visible(False)
buf7 = chart_img(fig)

s = prs.slides.add_slide(BLANK)
slide_header(s, '購物資料洞察', 'Shopping Intelligence Agent', accent_color=GREEN)
add_img(s, buf7, Inches(7.5), Inches(1.3), Inches(5.5), Inches(5.8))
points = [
    ('🎯 RFM 客戶分群',
     '依 Recency / Frequency / Monetary\n精準識別高價值客群，差異化行銷'),
    ('🔄 購買週期預測',
     '機器學習預測下次購買時機\n主動推送提醒，提升復購率 +23%'),
    ('🛍️ 商品關聯推薦',
     'Market Basket Analysis\n「買了A也買B」即時推薦，連帶銷售 +18%'),
    ('📅 節慶消費預測',
     '節日前4週啟動備貨與行銷\n庫存準確率提升 35%'),
]
for i, (title, body) in enumerate(points):
    y = Inches(1.45 + i * 1.45)
    add_rect(s, Inches(0.35), y, Inches(0.06), Inches(1.2), fill=GREEN)
    add_text(s, title, Inches(0.55), y, Inches(6.8), Inches(0.45),
             size=13, bold=True, color=DGRAY)
    add_text(s, body,  Inches(0.55), y + Inches(0.42), Inches(6.8), Inches(0.75),
             size=11.5, color=RGBColor(0x55,0x55,0x55))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Telecom Intelligence (with donut chart)
# ══════════════════════════════════════════════════════════════════════════
fig, ax = plt.subplots(figsize=(4.8, 4.0), facecolor='white')
labels = ['串流影音', '社群媒體', '一般上網', '通話/SMS', '遊戲']
sizes  = [32, 28, 20, 12, 8]
cols_d = ['#00B4D8','#06D6A0','#7B2FBE','#FF6B35','#FFC300']
wedges, texts, autotexts = ax.pie(
    sizes, labels=labels, colors=cols_d,
    autopct='%1.0f%%', startangle=90,
    wedgeprops=dict(width=0.55, edgecolor='white', linewidth=2),
    textprops={'fontsize': 9})
for at in autotexts:
    at.set_fontweight('bold')
ax.set_title('電信用戶流量行為分佈', fontsize=11, fontweight='bold')
buf8 = chart_img(fig)

s = prs.slides.add_slide(BLANK)
slide_header(s, '電信資料應用', 'Telecom Intelligence Agent', accent_color=BLUE)
add_img(s, buf8, Inches(7.8), Inches(1.3), Inches(5.0), Inches(5.5))
tpoints = [
    ('📉 客戶流失預警', '多維指標偵測流失信號\n預測準確率達 87%，提前30天介入'),
    ('🗺️ 生活型態分析', '上網行為 × 位置軌跡 × 消費時段\n建構「數位生活型態標籤」'),
    ('📍 商圈位置洞察', '基地台軌跡還原活動範圍\n商圈人流量 / 競品場域識別'),
    ('📊 ARPU 提升策略', '用量行為識別升級時機\n精準推送方案，ARPU +15%'),
]
for i, (title, body) in enumerate(tpoints):
    y = Inches(1.45 + i * 1.45)
    add_rect(s, Inches(0.35), y, Inches(0.06), Inches(1.2), fill=BLUE)
    add_text(s, title, Inches(0.55), y, Inches(7.0), Inches(0.45),
             size=13, bold=True, color=DGRAY)
    add_text(s, body,  Inches(0.55), y + Inches(0.42), Inches(7.0), Inches(0.75),
             size=11.5, color=RGBColor(0x55,0x55,0x55))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 9 — DCB Payment Intelligence
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, 'DCB 交易資料分析', 'DCB Payment Intelligence Agent', accent_color=ORANGE)

dcb_items = [
    ('💳 小額支付模式',  ORANGE, '分析電信代收交易頻率與金額\n識別高活躍消費用戶，LTV預測'),
    ('📦 訂閱服務分析',  ORANGE, '訂閱品項偏好 × 續約率分析\n主動推薦相關服務，ARPPU +20%'),
    ('💰 消費能力評估',  ORANGE, '支付行為建構替代信用模型\n服務無信用紀錄族群，普惠金融'),
    ('🚨 詐欺偵測引擎',  ORANGE, '即時異常交易識別（<100ms）\n規則引擎 + ML模型雙重防護'),
    ('🔗 跨平台整合',   ORANGE, 'DCB × 電商 × 實體通路串聯\n全渠道消費行為統一視圖'),
]
for i, (title, col, body) in enumerate(dcb_items):
    row, c_idx = divmod(i, 3)
    x = Inches(0.35 + c_idx * 4.3)
    y = Inches(1.4 + row * 2.55)
    add_rect(s, x, y, Inches(4.0), Inches(2.3),
             fill=WHITE, line=col, line_w=Pt(1.5))
    add_rect(s, x, y, Inches(4.0), Inches(0.5), fill=col)
    add_text(s, title, x + Inches(0.1), y + Inches(0.05),
             Inches(3.8), Inches(0.42), size=12, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, body, x + Inches(0.15), y + Inches(0.58),
             Inches(3.7), Inches(1.6), size=11.5, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 10 — Parking Intelligence
# ══════════════════════════════════════════════════════════════════════════
# heatmap-style chart
fig, ax = plt.subplots(figsize=(5.0, 4.0), facecolor='white')
hours = ['8am','10am','12pm','2pm','4pm','6pm','8pm','10pm']
zones = ['A商圈', 'B商圈', 'C商圈', 'D商圈']
data  = np.array([[60,75,90,70,65,85,95,50],
                  [40,50,80,85,70,90,88,45],
                  [30,35,55,60,80,75,70,40],
                  [20,25,40,45,55,60,50,30]])
im = ax.imshow(data, cmap='YlOrRd', aspect='auto', vmin=20, vmax=100)
ax.set_xticks(range(len(hours))); ax.set_xticklabels(hours, fontsize=9)
ax.set_yticks(range(len(zones))); ax.set_yticklabels(zones, fontsize=9)
ax.set_title('停車場人流熱力分佈 (停車次數/小時)', fontsize=10, fontweight='bold')
plt.colorbar(im, ax=ax, shrink=0.8, label='停車次數')
for i in range(len(zones)):
    for j in range(len(hours)):
        ax.text(j, i, str(data[i,j]), ha='center', va='center',
                fontsize=8, color='white' if data[i,j]>65 else 'black')
buf10 = chart_img(fig)

s = prs.slides.add_slide(BLANK)
slide_header(s, '停車資料商機', 'Parking Intelligence Agent', accent_color=PURPLE)
add_img(s, buf10, Inches(7.3), Inches(1.3), Inches(5.7), Inches(5.5))
pk = [
    ('📍 商圈熱力圖分析', '停車數據還原商圈人流密度\n識別黃金時段與冷熱區域'),
    ('🛒 停車×消費關聯', '停車後2小時內消費配對分析\n衍生「停車優惠券」導客入店'),
    ('🏪 零售選址決策', 'AI分析人流/競品/可及性\n選址準確率提升 40%'),
    ('📢 情境廣告投放', '進場瞬間精準推播附近優惠\n廣告點擊率提升 3.2倍'),
]
for i, (title, body) in enumerate(pk):
    y = Inches(1.45 + i * 1.45)
    add_rect(s, Inches(0.35), y, Inches(0.06), Inches(1.2), fill=PURPLE)
    add_text(s, title, Inches(0.55), y, Inches(6.5), Inches(0.45),
             size=13, bold=True, color=DGRAY)
    add_text(s, body, Inches(0.55), y + Inches(0.42), Inches(6.5), Inches(0.75),
             size=11.5, color=RGBColor(0x55,0x55,0x55))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 11 — Insurance Intelligence
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, '保險資料決策', 'Insurance Intelligence Agent', accent_color=GOLD)

ins_cards = [
    ('🎯 個人化風險評分', GOLD,
     '結合行為數據重新定義風險\n非傳統信用指標，服務更多客群'),
    ('📋 精準產品推薦', GREEN,
     '保單空白識別 + 需求預測\n交叉推薦成功率 +34%'),
    ('🚨 理賠詐欺偵測', ORANGE,
     '多源資料交叉比對\n異常理賠識別準確率 91%'),
    ('🔄 客戶生命週期', BLUE,
     '從首購→增購→挽留全程追蹤\nCLV 提升策略自動執行'),
    ('📈 動態費率定價', PURPLE,
     '行為數據驅動精算模型更新\n定價精準度提升 28%'),
    ('🤝 跨售機會', GOLD,
     '保險 × 電信 × 消費行為融合\n辨識最佳跨售時機與產品'),
]
for i, (title, col, body) in enumerate(ins_cards):
    row, c = divmod(i, 3)
    x = Inches(0.35 + c * 4.3)
    y = Inches(1.4 + row * 2.55)
    add_rect(s, x, y, Inches(4.0), Inches(2.3),
             fill=WHITE, line=col, line_w=Pt(1.5))
    add_rect(s, x, y, Inches(4.0), Inches(0.5), fill=col)
    add_text(s, title, x + Inches(0.1), y + Inches(0.05),
             Inches(3.8), Inches(0.42), size=12, bold=True,
             color=WHITE if col != GOLD else NAVY,
             align=PP_ALIGN.CENTER)
    add_text(s, body, x + Inches(0.15), y + Inches(0.58),
             Inches(3.7), Inches(1.6), size=11.5, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 12 — Cross-Domain Fusion
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
add_rect(s, 0, 0, W, Inches(0.95), fill=RGBColor(0x00, 0x1A, 0x3A))
add_text(s, '跨域資料融合洞察 | Cross-Domain Intelligence',
         Inches(0.5), Inches(0.1), Inches(12.3), Inches(0.78),
         size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

fusions = [
    ('🔵 + 🟣 + 🟢',
     '電信位置 × 停車軌跡 × 購物消費',
     '→ 精準商圈行銷',
     '識別消費者常駐商圈，\n在進停車場時推播個人化優惠券，\n轉化率提升 4.5倍',
     BLUE, '+52% 行銷效益'),
    ('🟠 + 🟡 + 🟢',
     'DCB支付 × 保險持有 × 購物行為',
     '→ 個人信用風險畫像',
     '建構替代信用評分，\n無卡族也能獲得金融服務資格，\n核批率提升 38%',
     ORANGE, '+38% 核批提升'),
    ('全域整合\n5大資料源',
     '購物 + 電信 + DCB + 停車 + 保險',
     '→ 360° 客戶智慧畫像',
     '9個Agent協同分析，\n實時更新客戶洞察標籤，\n個人化服務覆蓋率 95%',
     GREEN, '95% 客戶覆蓋'),
]
for i, (src, combo, result, detail, col, kpi) in enumerate(fusions):
    x = Inches(0.25 + i * 4.36)
    add_rect(s, x, Inches(1.1), Inches(4.1), Inches(5.8),
             fill=RGBColor(0x00, 0x22, 0x40), line=col, line_w=Pt(2))
    add_rect(s, x, Inches(1.1), Inches(4.1), Inches(0.55), fill=col)
    add_text(s, src,    x + Inches(0.1), Inches(1.12),
             Inches(3.9), Inches(0.5), size=13, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, combo,  x + Inches(0.1), Inches(1.72),
             Inches(3.9), Inches(0.55), size=11,
             color=RGBColor(0xAA, 0xCC, 0xFF), align=PP_ALIGN.CENTER)
    add_text(s, result, x + Inches(0.1), Inches(2.3),
             Inches(3.9), Inches(0.5), size=14, bold=True,
             color=col, align=PP_ALIGN.CENTER)
    add_text(s, detail, x + Inches(0.2), Inches(2.85),
             Inches(3.7), Inches(2.2), size=11.5,
             color=RGBColor(0xCC, 0xDD, 0xFF))
    add_rect(s, x + Inches(0.4), Inches(5.55), Inches(3.3), Inches(0.65),
             fill=col)
    add_text(s, '📊  ' + kpi, x + Inches(0.4), Inches(5.57),
             Inches(3.3), Inches(0.6), size=13, bold=True,
             color=NAVY if col == GREEN else WHITE,
             align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 13 — AI Decision Engine
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, 'AI 驅動決策引擎', 'AI-Powered Decision Engine — 3 Levels', accent_color=BLUE)

levels = [
    (Inches(0.35), Inches(1.4), LGRAY,   DGRAY, DGRAY,
     '📊 描述性分析 | Descriptive',
     '發生了什麼？ · What happened?',
     '即時 Dashboard · 自動報表 · KPI監控\nDatabricks SQL + Delta Live Tables',
     '回顧過去'),
    (Inches(0.35), Inches(3.1), RGBColor(0xE8,0xF5,0xFF), BLUE, BLUE,
     '🔮 預測性分析 | Predictive',
     '將會發生什麼？ · What will happen?',
     '客戶流失預測 · 需求預測 · 風險評分\nMLflow + Databricks AutoML + Feature Store',
     '洞察未來'),
    (Inches(0.35), Inches(4.8), RGBColor(0xEA,0xFB,0xF4), GREEN, GREEN,
     '🎯 規範性分析 | Prescriptive',
     '應該怎麼做？ · What should we do?',
     'Agent自動生成行動建議 · 即時干預執行\nMulti-Agent LLM + Databricks Workflows',
     '自動決策'),
]
for (x, y, bg_c, border, txt_col,
     title, sub, body, tag) in levels:
    add_rect(s, x, y, Inches(12.6), Inches(1.55),
             fill=bg_c, line=border, line_w=Pt(1.5))
    add_rect(s, x, y, Inches(1.6), Inches(1.55), fill=border)
    add_text(s, tag, x, y + Inches(0.5),
             Inches(1.6), Inches(0.55), size=12, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, title, x + Inches(1.75), y + Inches(0.05),
             Inches(6.0), Inches(0.5), size=14, bold=True, color=DGRAY)
    add_text(s, sub,   x + Inches(1.75), y + Inches(0.52),
             Inches(6.0), Inches(0.4), size=11, color=RGBColor(0x77,0x77,0x77))
    add_text(s, body,  x + Inches(1.75), y + Inches(0.9),
             Inches(10.5), Inches(0.6), size=11, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 14 — Business Use Cases
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, '商業應用場景', 'Business Use Cases — 6 High-Value Scenarios', accent_color=GREEN)

cases = [
    ('🎯 精準行銷',  GREEN,  '跨資料客群鎖定',
     '5源資料交叉建構精準受眾\n轉換率提升 45%，CPL降低 30%'),
    ('💳 信用評估',  BLUE,   '替代數據信用模型',
     'DCB+消費+電信建構信用分\n無卡族核批率 +38%'),
    ('🏪 智慧選址',  PURPLE, 'AI驅動展店決策',
     '停車+人流+競品分析選址\n選址成功率提升 40%'),
    ('💰 動態定價',  ORANGE, '即時需求感知定價',
     '時段×人流×庫存三角定價\n毛利率提升 12%'),
    ('🔄 客戶挽留',  GREEN,  '流失預警即時介入',
     '電信行為偵測30天前預警\n挽留成功率 67%'),
    ('🛡️ 風險控管', ORANGE, '多維度詐欺偵測',
     'DCB+保險+行為三層防護\n詐欺損失減少 42%'),
]
for i, (title, col, subtitle, body) in enumerate(cases):
    row, c = divmod(i, 3)
    x = Inches(0.35 + c * 4.3)
    y = Inches(1.4 + row * 2.7)
    add_rect(s, x, y, Inches(4.05), Inches(2.45),
             fill=WHITE, line=col, line_w=Pt(2))
    add_rect(s, x, y, Inches(4.05), Inches(0.55), fill=col)
    add_text(s, title, x + Inches(0.1), y + Inches(0.05),
             Inches(3.85), Inches(0.45), size=13, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, subtitle, x + Inches(0.1), y + Inches(0.6),
             Inches(3.85), Inches(0.38), size=11, bold=True,
             color=col, align=PP_ALIGN.CENTER)
    add_text(s, body, x + Inches(0.2), y + Inches(1.02),
             Inches(3.65), Inches(1.35), size=11, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 15 — ROI & Business Value (horizontal bar chart)
# ══════════════════════════════════════════════════════════════════════════
fig, ax = plt.subplots(figsize=(6.5, 4.5), facecolor='white')
metrics = ['行銷 ROI 提升', '客戶流失降低', '詐欺損失減少',
           '決策速度提升', '客戶終身價值', '選址成功率']
vals    = [35, 28, 42, 60, 22, 40]
cols_r  = ['#06D6A0','#00B4D8','#FF6B35','#7B2FBE','#FFC300','#06D6A0']
y_pos   = range(len(metrics))
bars = ax.barh(y_pos, vals, color=cols_r, edgecolor='white', linewidth=0.8, height=0.6)
for bar, v in zip(bars, vals):
    ax.text(bar.get_width() + 0.8, bar.get_y() + bar.get_height()/2,
            f'+{v}%', va='center', fontsize=11, fontweight='bold')
ax.set_yticks(y_pos); ax.set_yticklabels(metrics, fontsize=10)
ax.set_xlim(0, 75)
ax.set_xlabel('提升幅度 (%)', fontsize=10)
ax.set_title('Multi-Agent 平台商業價值量化', fontsize=12, fontweight='bold')
ax.grid(axis='x', alpha=0.3)
ax.spines['top'].set_visible(False); ax.spines['right'].set_visible(False)
buf15 = chart_img(fig)

s = prs.slides.add_slide(BLANK)
slide_header(s, '商業價值量化', 'Quantified Business Value & ROI', accent_color=GREEN)
add_img(s, buf15, Inches(0.3), Inches(1.2), Inches(7.0), Inches(6.0))
# KPI cards
kpis = [
    ('⚡', '決策速度', '提升 60%', BLUE),
    ('📉', '詐欺損失', '減少 42%', ORANGE),
    ('📈', '行銷 ROI', '提升 35%', GREEN),
    ('👥', '客戶 LTV',  '提升 22%', PURPLE),
]
for i, (icon, label, val, col) in enumerate(kpis):
    x = Inches(7.6)
    y = Inches(1.3 + i * 1.5)
    add_rect(s, x, y, Inches(5.4), Inches(1.35),
             fill=WHITE, line=col, line_w=Pt(2))
    add_rect(s, x, y, Inches(1.1), Inches(1.35), fill=col)
    add_text(s, icon, x, y + Inches(0.3),
             Inches(1.1), Inches(0.75), size=26,
             align=PP_ALIGN.CENTER, color=WHITE)
    add_text(s, label, x + Inches(1.2), y + Inches(0.12),
             Inches(4.0), Inches(0.45), size=13, color=DGRAY)
    add_text(s, val,   x + Inches(1.2), y + Inches(0.6),
             Inches(4.0), Inches(0.6), size=20, bold=True, color=col)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 16 — Implementation Roadmap
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
slide_header(s, '實施路徑圖', 'Implementation Roadmap — 4 Phases', accent_color=BLUE)

phases = [
    ('Phase 1\nM1–M2', '資料整合\n& Lakehouse建置',
     '• 資料盤點 & 品質評估\n• Unity Catalog架構設計\n• Delta Lake接入管線',
     BLUE),
    ('Phase 2\nM3–M4', 'Agent開發\n& 模型訓練',
     '• 9個Agent框架建置\n• ML特徵工程 & 訓練\n• POC場景驗證',
     GREEN),
    ('Phase 3\nM5–M6', '系統整合\n& 場景驗證',
     '• API整合 & 壓力測試\n• 業務場景上線驗收\n• 監控告警部署',
     ORANGE),
    ('Phase 4\nM7+', '規模化部署\n& 持續優化',
     '• 全量客戶上線\n• Agent自我學習優化\n• 新場景持續拓展',
     PURPLE),
]
# timeline bar
add_rect(s, Inches(0.35), Inches(2.25), Inches(12.6), Inches(0.12), fill=LGRAY)
for i, (ph, title, body, col) in enumerate(phases):
    x = Inches(0.35 + i * 3.18)
    # connector dot
    add_rect(s, x + Inches(1.2), Inches(2.1), Inches(0.35), Inches(0.35),
             fill=col)
    # card
    add_rect(s, x + Inches(0.05), Inches(2.55), Inches(3.05), Inches(4.55),
             fill=WHITE, line=col, line_w=Pt(2))
    add_rect(s, x + Inches(0.05), Inches(2.55), Inches(3.05), Inches(0.8), fill=col)
    add_text(s, ph,    x + Inches(0.1), Inches(1.3),
             Inches(2.95), Inches(0.8), size=11, bold=True,
             color=col, align=PP_ALIGN.CENTER)
    add_text(s, title, x + Inches(0.1), Inches(2.57),
             Inches(2.95), Inches(0.75), size=12, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, body,  x + Inches(0.2), Inches(3.45),
             Inches(2.75), Inches(3.5), size=11.5, color=DGRAY)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 17 — Technology Stack
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
add_rect(s, 0, 0, W, Inches(0.95), fill=RGBColor(0x00, 0x1A, 0x3A))
add_text(s, '技術架構棧 | Technology Stack',
         Inches(0.5), Inches(0.1), Inches(12.3), Inches(0.78),
         size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

layers = [
    ('🤖 AI / Agent Layer',
     'Claude API (Anthropic) · LangChain · CrewAI / AutoGen · MCP Protocol',
     BLUE, Inches(1.15)),
    ('📊 Analytics Layer',
     'Databricks MLflow · AutoML · Feature Store · Lakehouse AI',
     GREEN, Inches(2.3)),
    ('💾 Data Layer',
     'Delta Lake · Unity Catalog · Apache Spark · Delta Live Tables',
     ORANGE, Inches(3.45)),
    ('🔌 Integration Layer',
     'REST API · Apache Kafka · WebHook · Databricks Workflows · Dashboard',
     PURPLE, Inches(4.6)),
]
for label, tech, col, y in layers:
    add_rect(s, Inches(0.4), y, Inches(12.5), Inches(0.95),
             fill=RGBColor(0x00, 0x22, 0x40), line=col, line_w=Pt(1.5))
    add_rect(s, Inches(0.4), y, Inches(2.4), Inches(0.95), fill=col)
    add_text(s, label, Inches(0.45), y + Inches(0.2),
             Inches(2.3), Inches(0.55), size=11, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, tech, Inches(2.95), y + Inches(0.22),
             Inches(9.8), Inches(0.55), size=12,
             color=RGBColor(0xCC, 0xDD, 0xFF))

add_rect(s, Inches(0.4), Inches(5.75), Inches(12.5), Inches(0.55),
         fill=RGBColor(0x00, 0x33, 0x44))
add_text(s, '★  全棧部署於 Databricks Lakehouse — 統一資料、AI與分析的單一協作平台',
         Inches(0.5), Inches(5.8), Inches(12.3), Inches(0.45),
         size=13, color=GOLD, align=PP_ALIGN.CENTER)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 18 — Conclusion & Next Steps
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
add_rect(s, 0, 0, Inches(6.0), H, fill=RGBColor(0x00, 0x1A, 0x3A))
add_rect(s, Inches(0.45), Inches(0.6), Inches(0.1), Inches(5.5), fill=BLUE)

add_text(s, '結論與下一步', Inches(0.7), Inches(0.65),
         Inches(5.3), Inches(0.8), size=30, bold=True, color=WHITE)
add_text(s, 'Conclusion & Next Steps', Inches(0.7), Inches(1.42),
         Inches(5.3), Inches(0.45), size=14, color=BLUE)

takeaways = [
    '✅ Databricks 已具備打造 Multi-Agent 平台的完整資料基礎',
    '✅ 5大資料源融合創造傳統單源分析無法實現的洞察深度',
    '✅ 9個專業 Agent 協同可覆蓋從洞察到行動的全決策鏈路',
    '✅ 量化商業價值明確，ROI 可在 6個月內開始顯現',
]
for i, t in enumerate(takeaways):
    add_text(s, t, Inches(0.7), Inches(2.1 + i * 0.95),
             Inches(5.3), Inches(0.85), size=11.5,
             color=RGBColor(0xCC, 0xDD, 0xFF))

# right side — next steps
add_rect(s, Inches(6.4), Inches(0.5), Inches(6.5), Inches(6.5),
         fill=RGBColor(0x00, 0x22, 0x40), line=BLUE, line_w=Pt(1))
add_text(s, '⚡ 立即行動 | Immediate Actions',
         Inches(6.6), Inches(0.6), Inches(6.1), Inches(0.6),
         size=16, bold=True, color=BLUE)

actions = [
    ('STEP 1', '啟動資料盤點與品質評估',
     '盤點5大資料源覆蓋率、\n完整性與接入可行性', BLUE),
    ('STEP 2', '選定 2–3 個高價值 POC 場景',
     '建議優先：流失預警、精準行銷\n× DCB支付信用評估', GREEN),
    ('STEP 3', '組建跨職能 AI Agent 開發團隊',
     '資料工程師 + ML工程師 +\n業務分析師 + AI Agent開發者', ORANGE),
]
for i, (step, title, body, col) in enumerate(actions):
    y = Inches(1.35 + i * 1.68)
    add_rect(s, Inches(6.5), y, Inches(6.2), Inches(1.52),
             fill=RGBColor(0x00, 0x2A, 0x4A), line=col, line_w=Pt(1.5))
    add_rect(s, Inches(6.5), y, Inches(1.0), Inches(1.52), fill=col)
    add_text(s, step, Inches(6.5), y + Inches(0.5),
             Inches(1.0), Inches(0.52), size=11, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, title, Inches(7.6), y + Inches(0.08),
             Inches(4.9), Inches(0.5), size=13, bold=True,
             color=WHITE)
    add_text(s, body,  Inches(7.6), y + Inches(0.6),
             Inches(4.9), Inches(0.85), size=11,
             color=RGBColor(0xAA, 0xCC, 0xFF))

add_rect(s, Inches(6.4), Inches(6.45), Inches(6.5), Inches(0.6),
         fill=BLUE)
add_text(s, '讓資料說話，讓 AI 行動，讓商業持續進化。',
         Inches(6.5), Inches(6.5), Inches(6.3), Inches(0.5),
         size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ── Save ──────────────────────────────────────────────────────────────────
out = '/home/user/Moji/public/multi_agent_databricks_report.pptx'
prs.save(out)
print(f'Saved: {out}')
import os
print(f'Size: {os.path.getsize(out)/1024:.1f} KB  |  Slides: 18')
