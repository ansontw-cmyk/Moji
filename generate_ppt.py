from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree
import copy

# Colors
DARK_BLUE = RGBColor(0x0D, 0x1B, 0x3E)
MID_BLUE  = RGBColor(0x1A, 0x2F, 0x6B)
ACCENT_BLUE = RGBColor(0x2E, 0x6D, 0xCC)
LIGHT_BLUE = RGBColor(0x5D, 0xAD, 0xE2)
RED       = RGBColor(0xC0, 0x39, 0x2B)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_TEXT = RGBColor(0xCC, 0xD6, 0xF1)
GOLD      = RGBColor(0xF3, 0x9C, 0x12)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

slide_layout = prs.slide_layouts[6]  # blank
slide = prs.slides.add_slide(slide_layout)

def add_rect(slide, x, y, w, h, fill_rgb=None, line_rgb=None, line_w=Pt(0)):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.width = line_w
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb:
        shape.line.color.rgb = line_rgb
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, x, y, w, h, size=Pt(14), bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(x, y, w, h)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox

# ── Background ───────────────────────────────────────────────────────────────
bg = add_rect(slide, 0, 0, W, H, fill_rgb=DARK_BLUE)

# subtle grid overlay (thin horizontal lines)
for i in range(1, 8):
    line_shape = add_rect(slide, 0, Inches(i), W, Pt(0.4),
                          fill_rgb=RGBColor(0x1E, 0x35, 0x6E))

# ── Title block ──────────────────────────────────────────────────────────────
add_text(slide, "AI-Driven SDLC",
         Inches(0.4), Inches(0.18), Inches(7), Inches(0.65),
         size=Pt(32), bold=True, color=WHITE)

add_text(slide, "新一代 AI 驅動的軟體開發生命週期",
         Inches(0.4), Inches(0.78), Inches(8), Inches(0.45),
         size=Pt(17), bold=False, color=LIGHT_BLUE)

# ── Company logo placeholder ──────────────────────────────────────────────────
add_rect(slide, Inches(11.8), Inches(0.18), Inches(1.3), Inches(0.5),
         fill_rgb=RGBColor(0x1A, 0x2F, 0x6B),
         line_rgb=LIGHT_BLUE, line_w=Pt(1))
add_text(slide, "遠傳 FET", Inches(11.8), Inches(0.22), Inches(1.3), Inches(0.42),
         size=Pt(11), bold=True, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

# ── AI Gatekeeper banner ──────────────────────────────────────────────────────
gk_y = Inches(1.38)
add_rect(slide, Inches(0.35), gk_y, Inches(12.63), Pt(2),
         fill_rgb=RGBColor(0x2E, 0x6D, 0xCC))

add_text(slide, "AI Gatekeeper  ─── 全程把關每個 SDLC 節點",
         Inches(0.35), Inches(1.32), Inches(12.63), Inches(0.42),
         size=Pt(13), bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ── Stage definitions ─────────────────────────────────────────────────────────
stages = [
    ("01", "需求",  "AI 起草",    "人類決策"),
    ("02", "設計",  "AI 提案",    "架構師定案"),
    ("03", "開發",  "Human + AI", "結對開發"),
    ("04", "測試",  "AI 全自動",  "測試覆蓋"),
    ("05", "文件",  "AI 即時",    "雙語文件化"),
]

icons = ["📋", "📐", "</>", "🧪", "📄"]

box_w   = Inches(2.3)
box_h   = Inches(3.6)
box_y   = Inches(1.6)
gap     = Inches(0.18)
start_x = Inches(0.35)

for i, (num, title, line1, line2) in enumerate(stages):
    bx = start_x + i * (box_w + gap)

    # Card background
    card = add_rect(slide, bx, box_y, box_w, box_h,
                    fill_rgb=MID_BLUE,
                    line_rgb=ACCENT_BLUE, line_w=Pt(1.2))

    # Number badge
    badge = add_rect(slide, bx + Inches(0.12), box_y + Inches(0.12),
                     Inches(0.52), Inches(0.38),
                     fill_rgb=ACCENT_BLUE)
    add_text(slide, num,
             bx + Inches(0.12), box_y + Inches(0.1),
             Inches(0.52), Inches(0.42),
             size=Pt(13), bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Stage title
    add_text(slide, title,
             bx, box_y + Inches(0.55),
             box_w, Inches(0.48),
             size=Pt(20), bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Icon area (emoji as stand-in)
    add_text(slide, icons[i],
             bx, box_y + Inches(1.1),
             box_w, Inches(0.9),
             size=Pt(36), bold=False, color=WHITE, align=PP_ALIGN.CENTER)

    # Divider
    add_rect(slide, bx + Inches(0.2), box_y + Inches(2.1),
             box_w - Inches(0.4), Pt(1.5),
             fill_rgb=ACCENT_BLUE)

    # AI role lines (red accent)
    add_text(slide, line1,
             bx, box_y + Inches(2.22),
             box_w, Inches(0.38),
             size=Pt(13), bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_text(slide, line2,
             bx, box_y + Inches(2.6),
             box_w, Inches(0.38),
             size=Pt(13), bold=False, color=GRAY_TEXT, align=PP_ALIGN.CENTER)

# Arrows between cards
arrow_y = box_y + Inches(1.65)
for i in range(4):
    ax = start_x + (i + 1) * (box_w + gap) - gap
    add_text(slide, "▶",
             ax - Inches(0.02), arrow_y,
             gap + Inches(0.04), Inches(0.35),
             size=Pt(14), bold=False, color=LIGHT_BLUE, align=PP_ALIGN.CENTER)

# ── Bottom Line ───────────────────────────────────────────────────────────────
bl_y = Inches(5.42)
add_rect(slide, Inches(0.35), bl_y, Inches(12.63), Inches(1.72),
         fill_rgb=RGBColor(0x0A, 0x14, 0x2E),
         line_rgb=RED, line_w=Pt(2))

# Red left bar
add_rect(slide, Inches(0.35), bl_y, Inches(0.1), Inches(1.72),
         fill_rgb=RED)

add_text(slide, "Bottom Line",
         Inches(0.55), bl_y + Inches(0.12),
         Inches(3), Inches(0.38),
         size=Pt(12), bold=True, color=RED)

add_text(slide, "AI Gatekeeper 全程把關",
         Inches(0.55), bl_y + Inches(0.5),
         Inches(12.2), Inches(0.45),
         size=Pt(22), bold=True, color=WHITE)

add_text(slide, "每個節點自動過濾重複工作，工程師專注高價值決策",
         Inches(0.55), bl_y + Inches(0.98),
         Inches(12.2), Inches(0.42),
         size=Pt(14), bold=False, color=LIGHT_BLUE)

# ── Save ──────────────────────────────────────────────────────────────────────
out = "/home/user/Moji/AI_SDLC_beautified.pptx"
prs.save(out)
print(f"Saved: {out}")
